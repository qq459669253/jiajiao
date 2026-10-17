"""GPL-3.0. Expanded, self-contained teaching explanation for the capstone."""
import io
import math
import zipfile
from pathlib import Path

import fitz
from docx.shared import Cm, Pt
from reportlab.pdfgen import canvas

import build_selected_pack as base
from illustrated_teacher import Drawing, add, sub, mul, mid, length

MM = base.MM
RED = (.72, .12, .14)
BLUE = (.05, .32, .66)
GREEN = (.08, .43, .31)
GRAY = (.55, .58, .61)


def step(group, title, key, why, action, proof, result, check):
    return dict(group=group, title=title, key=key, why=why, action=action,
                proof=proof, result=result, check=check)


LESSONS = [
    step('先备知识', '引理第1步：把中点条件变成全等条件', 'lemma_sas',
         '后三问都要用“连接两边中点的线段平行于第三边”。若尚未学过，先用全等证明；本页和下一页合起来就是完整证明。',
         '在△TUV中，X、Y分别为TU、TV中点。延长XY过Y至Z，使YZ=YX；连接VZ、UZ。',
         '① TY=YV（Y是TV中点），YX=YZ（作图）。\n② ∠TYX=∠VYZ（对顶角）。\n③ 所以△TYX≌△VYZ（SAS）。\n④ 对应边相等，得TX=VZ；又TX=XU，故XU=VZ。\n⑤ 对应角∠TXY=∠VZY；它们是内错角，故TX∥VZ。\n⑥ T、X、U共线，因此XU∥VZ。',
         '得到一组平行且相等的线段：XU∥VZ，XU=VZ。',
         '本页只做准备：还没有证明XY∥UV，要继续看第二次全等。'),
    step('先备知识', '引理第2步：再用一次SAS，得到平行', 'lemma_parallel',
         '上页已把XU与VZ变成平行且相等的一组边。连接UZ后，再比较共用UZ的两个三角形。',
         '突出△XUZ、△VZU；UZ是公共边。',
         '① XU=VZ（上页），UZ=ZU（公共边）。\n② XU∥VZ，故∠XUZ=∠VZU（内错角）。\n③ 所以△XUZ≌△VZU（SAS）。\n④ 得∠XZU=∠VUZ（对应角）。\n⑤ 这两个角是直线XZ、UV被UZ所截的内错角，故XZ∥UV。\n⑥ X、Y、Z共线，所以XY∥UV。',
         '三角形两边中点的连线平行于第三边；下文简称“中点连线性质”。',
         '对应到后文：先找到一个三角形，再说明两个点分别是哪两条边的中点。'),
    step('第（1）问', '第1步：为什么要倍长BF？', 'h1_sas',
         'AB=CD，但这两条边还没在同一个三角形里。想把AB复制到D处；F又是AD中点，所以尝试以F为中心作一组全等三角形。',
         '连接BF，延长BF过F至H，使FH=FB；连接DH、CH。',
         '① AF=DF（F是AD中点），BF=FH（作图）。\n② A、F、D共线，B、F、H共线，故∠AFB=∠DFH。\n③ △AFB≌△DFH（SAS），对应为A↔D、F↔F、B↔H。\n④ 所以DH=AB；结合AB=CD，得DH=CD。\n⑤ 还得到∠ABF=∠DHF（对应角）。\n⑥ 这是一组内错角，故AB∥DH。',
         'DH=CD，AB∥DH；同时B、F、H共线且BF=FH。',
         '辅助点H做了两件事：复制AB，并让F成为BH的中点。'),
    step('第（1）问', '第2步：两个中点落在哪个三角形里？', 'h1_mid',
         'E原本就是BC中点；上一步又得到F是BH中点。两条线段BC、BH共同属于△BHC，所以应当看这个大三角形。',
         '突出△BHC及其中的EF；另用蓝色突出等腰△DHC。',
         '① 在△BHC中，BE=EC，故E是BC中点。\n② B、F、H共线且BF=FH，故F是BH中点。\n③ 由中点连线性质，EF∥CH。\n④ 由上页DH=CD，△DHC是等腰三角形，底边为HC。\n⑤ 因而两个底角相等：∠DHC=∠DCH。\n⑥ 现在已知AB∥DH、EF∥CH，下一步把目标角对应过来。',
         '先在等腰△DHC中得到∠DHC=∠DCH，再去找原题的角。',
         '等腰三角形的顶点是D，两个相等的角在H、C。'),
    step('第（1）问', '第3步：把两组对应角分别画出来', 'h1_angles',
         '“两组线分别平行”本身还不够，角也可能互补。要检查组成角的两条射线是同向还是反向；两条射线同时反向，角度保持不变。',
         '上排对照M、H处的角；下排对照N、C处的角。箭头表示射线方向，同色弧表示待比较的角。',
         '① 在图a的位置中，射线MB与HD平行且同向，ME与HC也平行且同向。\n   所以∠BME=∠DHC（上排两个角）。\n② 射线NC与CD平行且反向，NE与CH也平行且反向。\n   相当于把整个角转过180°，所以∠CNE=∠DCH（下排两个角）。\n③ 上页已证∠DHC=∠DCH。\n④ 连起来：∠BME=∠DHC=∠DCH=∠CNE。',
         '∠BME=∠CNE，第（1）问得证。',
         '不能把“CD与NC共线”直接写成同向；这里它们实际反向。'),
    step('第（2）问', '第1步：位置变了，哪些结论仍成立？', 'h2_sas',
         '第（1）问的构造只用到中点和全等，换成图b仍可使用。先重新列出理由，再辨认目标三角形的位置。',
         '仍延长BF过F至H，使FH=FB；连接DH、CH。',
         '① AF=DF、BF=FH，且∠AFB=∠DFH。\n② △AFB≌△DFH（SAS），得DH=AB=CD。\n③ ∠ABF=∠DHF（对应角），内错角相等，故AB∥DH。\n④ 在△BHC中，E是BC中点，F是BH中点，故EF∥CH。\n⑤ M、N、E、F共线，所以MN∥CH。\n⑥ DH=CD，所以∠DHC=∠DCH。',
         '仍有等腰△DHC、MN∥CH、AB∥DH。',
         'O是AB与CD交点；M、N是EF分别与CD、AB的交点。不要把交点认反。'),
    step('第（2）问', '第2步：小三角形的底角对应哪里？', 'h2_angles',
         '原图的△OMN很小，容易看错。这里把M、N处的角拆开画，并与等腰△DHC的两个底角分别比较。',
         '上排比较∠OMN与∠DCH；下排比较∠ONM与∠DHC。',
         '① 射线MO与CD平行且反向，MN与CH也平行且反向。\n   两条射线同时反向，故∠OMN=∠DCH。\n② 射线NO与HD平行且反向，NM与HC也平行且反向。\n   两条射线同时反向，故∠ONM=∠DHC。\n③ 等腰△DHC给出∠DCH=∠DHC。\n④ 所以∠OMN=∠ONM，即△OMN的M、N两角相等。',
         '相等的是M、N处的两个角；下一页由等角判断等边。',
         '这些射线方向按图b所示位置判断；角的拆图保留了原图的方向。'),
    step('第（2）问', '第3步：由两个相等角，写出形状', 'h2_result',
         '已经得到M、N处角相等，形状判断只需再用“等角对等边”。此时应单独看△OMN。',
         '放大△OMN，标出相等的底角及其所对的边。',
         '① ∠OMN=∠ONM（上页）。\n② ∠OMN所对的边是ON，∠ONM所对的边是OM。\n③ 在同一个三角形中，等角对等边，所以ON=OM。\n④ 故△OMN是以O为顶点、MN为底边的等腰三角形。',
         '△OMN是等腰三角形，OM=ON。',
         '条件不能确定O处角度；这里没有理由进一步断言它是直角三角形。'),
    step('第（3）问', '第1步：辅助点T是怎样反推出来的？', 'h3_construct',
         'E是BC中点。若想让EF平行于某条新边TB，就应让F成为TC的中点。因此先反推：怎样取T才能使TF=FC？',
         '延长CA过A至T，使AT=CD（也等于AB）；本页先只看这条直线上的长度。',
         '① A、F、D、C依次共线，AF=FD（F是AD中点）。\n② 若T放在A外侧，则TF=TA+AF，而FC=FD+DC。\n③ 因为AF=FD，只要取TA=DC，就能使上述两段相等。\n④ 因此作T，使AT=CD。再用CD=AB，得AT=AB。\n⑤ 此时T、A、F、D、C依次共线，\n   TF=TA+AF=DC+FD=FC，所以F是TC的中点。',
         '取T的目的：把“F是AD中点”变成“F也是TC中点”；同时得到AT=AB。',
         'T不是凭空猜出来的：先想要TF=FC，再倒推应取TA=DC。'),
    step('第（3）问', '第2步：60°先使哪个三角形等边？', 'h3_tab',
         '现在E、F已成为△TBC两边的中点。先得到EF∥TB，把已知60°转移到T处；AT=AB便能接着发挥作用。',
         '连接TB，突出△TAB，标出T处与F处对应的60°。',
         '① 在△TBC中，E是BC中点，F是TC中点，故EF∥TB。\n② 直线TC为截线，∠BTA与∠EFC为同位角。\n   因而∠BTA=∠EFC=60°。\n③ AT=AB，故△TAB中∠TBA=∠BTA=60°。\n④ 由三角形内角和，∠TAB=180°-60°-60°=60°。\n⑤ 三个角都为60°，所以△TAB是等边三角形。',
         '△TAB等边，∠TAB=∠BTA=∠TBA=60°，且EF∥TB。',
         '已知的60°是∠EFC；这里必须先证EF∥TB，才可以把它搬到T。'),
    step('第（3）问', '第3步：A处的60°怎样得到？', 'h3_at_a',
         '最终要看△AGD，而△TAB在另一侧。先利用A点两条交叉的直线，把等边△TAB的60°转到△AGF中。',
         '单独画出G—A—B与T—A—F两条相交直线；同色弧标出一对对顶角。',
         '① G在BA越过A的延长线上，所以G、A、B依次共线。\n② T在CA越过A的延长线上，且F在AC上，所以T、A、F依次共线。\n③ 因而射线AG与AB反向，射线AF与AT反向。\n④ ∠GAF与∠TAB是一对对顶角，大小相等。\n⑤ 上页得到∠TAB=60°，故∠GAF=60°。',
         '△AGF在A处的角为60°；还需再确定它的另一个角。',
         '请先在图上指出组成∠TAB的两条射线，再找各自的反向射线。'),
    step('第（3）问', '第4步：F处再得60°，得到第二个等边三角形', 'h3_at_f',
         '只知道一个60°还不能认定等边。利用GF所在直线与TB平行，再得到F处的60°，才能推出GF=AF。',
         '突出平行的TB、FG，以及截线TC；比较T处和F处的一组内错角。',
         '① G、F、E共线，且EF∥TB，所以FG∥TB。\n② 直线TC截平行线TB、FG，∠BTA与∠AFG是一组内错角。\n   因而∠AFG=∠BTA=60°。\n③ 上页已得∠GAF=60°，所以∠AGF=180°-60°-60°=60°。\n④ △AGF三个角均为60°，故为等边三角形，AG=GF=AF。\n⑤ 原题F是AD中点，AF=FD，所以GF=FD。',
         '△AGF等边；尤其GF=AF=FD，这是下一步判断等腰的依据。',
         'F处的60°是∠AFG；相邻的∠GFD还没有算，它不是60°。'),
    step('第（3）问', '第5步：为什么会出现120°和30°？', 'h3_gfd',
         'F在AD上，FA与FD方向相反。因此等边△AGF的60°旁边，恰好是一个120°；再利用GF=FD求出两个底角。',
         '单独放大△GFD，保留A以便看清FA、FD是两条反向射线。',
         '① A、F、D依次共线，故∠GFA+∠GFD=180°（平角）。\n② △AGF等边，∠GFA=60°，所以∠GFD=120°。\n③ 上页GF=FD，故△GFD是等腰三角形，顶点为F。\n④ ∠FGD=∠FDG=(180°-120°)÷2=30°。\n⑤ A、F均在D的同侧，且三点共线，所以射线DF与DA重合，\n   因而∠ADG=∠FDG=30°。',
         '△GFD中：F处120°，G、D处各30°；△AGD在D处也是30°。',
         '相等边GF、FD的公共端点是F，所以120°是顶角，两个30°是底角。'),
    step('第（3）问', '第6步：哪个角是直角？把它拼出来', 'h3_final',
         '要求判断的是△AGD。前面得到的等边△AGF、等腰△GFD只是中间结构，最后必须回到G点把相邻的两个角相加。',
         '突出整个△AGD，用GF把G处的角分成60°与30°两部分，再标出总直角。',
         '① 在△AGF中，∠AGF=60°；在△GFD中，∠FGD=30°。\n② F在线段AD上，所以射线GF在∠AGD内部。\n③ ∠AGD=∠AGF+∠FGD=60°+30°=90°。\n④ 因而△AGD是直角三角形，直角顶点为G。\n⑤ 另有∠ADG=30°、∠GAD=60°。\n⑥ 还可由AD=AF+FD=2AF=2AG，得到AD=2AG。',
         '最终答案：△AGD是直角三角形，∠AGD=90°，∠ADG=30°。',
         '不要把两个中间结论混起来：△AGF等边、△GFD等腰、△AGD直角。'),
]


def arrow(c, p, q, color, width=1.25):
    c.saveState()
    c.setStrokeColorRGB(*color)
    c.setLineWidth(width)
    c.line(*p, *q)
    v = mul(sub(q, p), 1 / length(sub(q, p)))
    n = (-v[1], v[0])
    for sign in [-1, 1]:
        z = add(sub(q, mul(v, 5)), mul(n, sign * 2.4))
        c.line(*q, *z)
    c.restoreState()


def angle_comparisons(c, b, kind):
    p = b.geometry(kind)
    pairs = ([('M', 'B', 'E'), ('H', 'D', 'C'),
              ('N', 'C', 'E'), ('C', 'D', 'H')] if kind == 'h1' else
             [('M', 'O', 'N'), ('C', 'D', 'H'),
              ('N', 'O', 'M'), ('H', 'D', 'C')])
    for i, (v, a, z) in enumerate(pairs):
        color = RED if i < 2 else BLUE
        directions = [mul(sub(p[name], p[v]), 49 / length(sub(p[name], p[v]))) for name in [a,z]]
        # Give each angle its own bounding box. Upward and downward rays
        # have different vertex heights, so titles and endpoints never collide.
        ends = [mul(u,1.17) for u in directions]+[(0,0)]
        lo,hi = (125,185) if i<2 else (15,75)
        center = ((25+143*(i%2)+116+143*(i%2)-min(x for x,y in ends)-max(x for x,y in ends))/2,
                  (lo+hi-min(y for x,y in ends)-max(y for x,y in ends))/2)
        for name in [a, z]:
            u = sub(p[name], p[v]); u = mul(u, 49 / length(u))
            end = add(center, u)
            arrow(c, center, end, color)
            loc = add(center, mul(u, 1.17))
            base.txt(c, name, loc[0] - 3, loc[1] - 3, width=20, size=9, lead=10)
        u, w = directions
        start = math.degrees(math.atan2(u[1], u[0]))
        end = math.degrees(math.atan2(w[1], w[0]))
        extent = (end - start + 180) % 360 - 180
        c.setStrokeColorRGB(*color)
        c.arc(center[0]-17, center[1]-17, center[0]+17, center[1]+17, start, extent)
        away=mul(add(u,w),-1/length(add(u,w)))
        vl=add(center,mul(away,10))
        base.txt(c, v, vl[0]-3, vl[1]-3, width=20, size=9, lead=10)
        base.txt(c, f'∠{a}{v}{z}', 20+143*(i % 2), 204-109*(i//2), width=100, size=10, lead=12)
    c.setStrokeColorRGB(*GRAY);c.setLineWidth(.5);c.line(8,110,252,110)


def figure(c, key, b):
    if key in ['h1_angles', 'h2_angles']:
        angle_comparisons(c, b, key[:2]);return
    if key.startswith('lemma'):
        p = {'T':(0,3), 'U':(-3,0), 'V':(3,0), 'X':(-1.5,1.5), 'Y':(1.5,1.5), 'Z':(4.5,1.5)}
        d = Drawing(c, p, base.txt)
        d.edges([('T','U'),('U','V'),('V','T'),('X','Z'),('V','Z'),('U','Z')])
        if key == 'lemma_sas':
            d.triangle(['T','Y','X'], RED);d.triangle(['V','Y','Z'], BLUE)
            for a,z in [('T','Y'),('Y','V')]:d.tick(a,z)
            for a,z in [('X','Y'),('Y','Z')]:d.tick(a,z,2)
        else:
            d.triangle(['X','U','Z'], RED);d.triangle(['V','Z','U'], BLUE)
            d.tick('X','U');d.tick('V','Z');d.segment('U','Z',GREEN,2)
            d.segment('X','Y',RED,2);d.segment('U','V',BLUE,2)
        d.labels({'T':(-4,8),'U':(-13,-12),'V':(3,-12),'X':(-15,7),'Y':(-4,8),'Z':(4,4)})
        return
    if key in ['h1_sas', 'h2_sas', 'h2_result', 'm1', 'm2']:
        b.diagram(c,key);return
    p = b.geometry(key.split('_')[0])
    if key == 'h1_mid':
        d = Drawing(c,p,base.txt)
        d.edges([('A','B'),('B','C'),('C','D'),('D','A'),('B','H'),('D','H'),('C','H'),('M','E'),('N','E')])
        d.triangle(['B','H','C'],GREEN);d.triangle(['D','H','C'],BLUE)
        d.segment('E','F',RED,2)
        for a,z in [('B','E'),('E','C')]:d.tick(a,z,2)
        for a,z in [('B','F'),('F','H')]:d.tick(a,z)
        d.tick('D','H',3);d.tick('D','C',3)
        d.labels({'M':(-14,5),'N':(-16,4),'A':(-15,0),'B':(-14,-12),'C':(5,-10),'D':(6,-4),'E':(-3,-15),'F':(-13,6),'H':(5,5)})
        return
    if key == 'h3_construct':
        # Length diagram, with four deliberately separated subsegments.
        d = Drawing(c,{k:p[k] for k in ['T','A','F','D','C']},base.txt)
        d.segment('T','C',GRAY,1.2)
        d.segment('T','A',BLUE,2);d.segment('D','C',BLUE,2)
        d.segment('A','F',RED,2);d.segment('F','D',RED,2)
        d.tick('T','A',2);d.tick('D','C',2)
        d.tick('A','F');d.tick('F','D')
        d.labels({k:(-3,10 if k in ['A','D'] else -16) for k in p})
        base.txt(c,'蓝色：TA=DC；红色：AF=FD',28,40,width=230,size=10,lead=13)
        return
    keep = {'h3_tab':['T','A','B','C','E','F','G'],
            'h3_at_a':['T','A','B','F','G'],
            'h3_at_f':['T','A','B','F','G'],
            'h3_gfd':['A','F','D','G'],
            'h3_final':['A','F','D','G']}[key]
    p = {k:p[k] for k in keep};d=Drawing(c,p,base.txt)
    offsets={'T':(-15,5),'A':(-14,-3),'B':(-15,-12),'C':(4,-12),'E':(-3,-15),'F':(6,-8),'G':(-3,8),'D':(5,-8)}
    if key == 'h3_tab':
        d.edges([('T','C'),('B','C'),('A','B'),('G','B'),('G','E')])
        d.triangle(['T','A','B'],BLUE)
        d.segment('T','B',RED,1.7);d.segment('F','E',RED,1.7)
        d.tick('T','A');d.tick('A','B')
        d.angle('T','B','A','60°',RED,16);d.angle('F','E','C','60°',RED,16)
    elif key == 'h3_at_a':
        d.edges([('T','F'),('G','B')]);d.segment('T','B',GRAY,.8,True)
        d.angle('A','T','B','60°',BLUE,17);d.angle('A','G','F','60°',RED,17)
    elif key == 'h3_at_f':
        d.edges([('T','F'),('G','B')]);d.segment('T','B',BLUE,1.7);d.segment('F','G',RED,1.7)
        d.triangle(['A','G','F'],RED)
        d.angle('T','B','A','60°',BLUE,16);d.angle('F','A','G','60°',RED,16)
        d.tick('A','G');d.tick('G','F');d.tick('A','F')
    else:
        d.edges([('A','D'),('A','G'),('G','D'),('G','F')])
        if key == 'h3_gfd':
            d.triangle(['G','F','D'],BLUE);d.tick('G','F');d.tick('F','D')
            d.angle('F','G','D','120°',BLUE,20)
            d.angle('G','F','D','30°',BLUE,21);d.angle('D','F','G','30°',BLUE,26)
        else:
            d.triangle(['A','G','D'],RED);d.segment('G','F',BLUE,1.5)
            d.angle('G','A','F','60°',RED,17);d.angle('G','F','D','30°',BLUE,33)
            d.right('G','A','D',GREEN)
    d.labels(offsets)


def render_figures(b):
    for s in LESSONS:
        buf=io.BytesIO();c=canvas.Canvas(buf,pagesize=(260,220));figure(c,s['key'],b);c.save()
        pdf=buf.getvalue();path=b.ROOT/'teaching_figures'/s['key']
        path.with_suffix('.pdf').write_bytes(pdf)
        doc=fitz.open(stream=pdf,filetype='pdf')
        doc[0].get_pixmap(matrix=fitz.Matrix(3,3)).save(path.with_suffix('.png'))


def heading(c,subtitle,page):
    base.header(c,'全等与中点：逐步配图教师详解',subtitle,page,revision='2026-10-17 · 详解修订版')


def lesson_page(c,s,page,b):
    heading(c,s['group']+'｜一页一个关键结论',page)
    y=base.txt(c,s['title'],18*MM,254*MM,size=12.5,lead=18)
    y=base.txt(c,'为什么这样想：'+s['why'],18*MM,y-6,size=10.5,lead=16)
    if y<225*MM:raise ValueError('Motivation overflow: '+s['key'])
    c.saveState();c.translate(49*MM,132*MM);c.scale(112*MM/260,94.77*MM/220);figure(c,s['key'],b);c.restoreState()
    y=base.txt(c,'看图／作图：'+s['action'],18*MM,126*MM,size=9.5,lead=14)
    y=base.txt(c,s['proof'],18*MM,min(y-7,111*MM),size=10.5,lead=17)
    if y<43*MM:raise ValueError('Proof overflow: '+s['key'])
    c.setStrokeColorRGB(*BLUE);c.setLineWidth(.6);c.line(18*MM,42*MM,192*MM,42*MM)
    y=base.txt(c,'本页结论：'+s['result'],18*MM,35*MM,size=10.5,lead=16)
    y=base.txt(c,'讲解提醒：'+s['check'],18*MM,y-4,size=9.5,lead=14)
    if y<16*MM:raise ValueError('Conclusion overflow: '+s['key'])
    c.showPage()


def overview(c,page,b):
    heading(c,'三问压轴题｜先核对题面与读图顺序',page)
    y=base.txt(c,'本版怎样阅读',18*MM,252*MM,size=12,lead=18)
    y=base.txt(c,'每一步都分成“为什么这样想—图上怎样看—逐条证明—本页结论”。先备引理可按学生进度选择讲解；三问均给出完整理由，不默认孩子会从复杂图中自行认出对应角。',18*MM,y-5,size=10.5,lead=16)
    y=base.txt(c,'原题（原图所示位置）',18*MM,y-10,size=11,lead=17)
    y=base.txt(c,b.HARD,18*MM,y-4,size=10,lead=17)
    b.image_fit(c,b.ROOT/'sources/challenge_630.png',18*MM,86*MM,174*MM,52*MM)
    base.txt(c,'读图先记住：\n图a：M、N是延长线上的交点；图b：O=AB与CD的交点。\n图c：G—A—B共线，A—F—D—C共线；T是解题时新作的点。\n本版保留原题，只重写教师解释。图b按原图补充的交点定义及出处见README。',18*MM,76*MM,size=10,lead=17)
    c.showPage()


RECAP = [
    ('第（1）问：从全等到等角',
     '倍长BF至H → △AFB≌△DFH → DH=CD且AB∥DH → 在△BHC中得EF∥CH → △DHC等腰 → 分别对应M、N处的角 → ∠BME=∠CNE。'),
    ('第（2）问：同一构造，重新辨认角',
     '仍倍长BF至H → DH=CD、AB∥DH、MN∥CH → ∠OMN=∠DCH、∠ONM=∠DHC → M、N两角相等 → OM=ON。只能确定等腰。'),
    ('第（3）问：三个三角形，各有不同作用',
     '反推TF=FC，取TA=DC → EF∥TB → △TAB等边 → 对顶角和内错角使△AGF等边 → GF=AF=FD → △GFD等腰，得120°及两个30° → 在G处相加60°+30°，所以△AGD直角。'),
    ('教师追问及可核对的回答',
     '问：为什么T在A外侧？答：这样TF=TA+AF，能与FC=FD+DC逐项比较，取TA=DC即可。\n问：为何不能见60°就认等边？答：必须再用等腰条件或求出第二个60°。\n问：GF=FD从哪来？答：△AGF等边给GF=AF，原中点给AF=FD。\n问：直角在哪里？答：G处，因为∠AGD=∠AGF+∠FGD=60°+30°。'),
    ('讲课时的用法',
     '先让学生在原图上指出对应的三角形、边、角，再展示下一页。讲后遮住详解，要求学生重画H或T并复述每一个等号的理由。已经会的页可以跳过，无须把整份详解当作逐字朗读稿。'),
]


def recap(c,page):
    heading(c,'讲完后的复盘｜用同一题检查是否真正理解',page)
    y=252*MM
    for title,body in RECAP:
        y=base.txt(c,title,18*MM,y,size=11.5,lead=18)
        y=base.txt(c,body,18*MM,y-5,size=10.5,lead=17)-13
    if y<22*MM:raise ValueError('Recap overflow')
    c.showPage()


INTRO = [
    ('本次修订了什么', '原版把前两问的对应角、第三问的辅助点动机写得过快。本版把三问拆成12个图文步骤，并增加两页中点引理；一个关键结论占一页，射线方向和角的来源单独解释。学生题面、选题和分值保持一致。'),
    ('两题的定位与安排', '题1是中等题，建议10—12分钟；题2是综合压轴题，建议25—35分钟。两题独立完成约40—45分钟，再用约45分钟讲解，余下时间用于批改、追问和遮住答案后重建辅助线。'),
    ('先备知识', '全等SAS、平行线与角、等腰和等边三角形、中点。若中点连线性质尚未学，先看本版两页引理；证明只用两次SAS，无需相似、勾股定理或根式。'),
    ('答案与评分', '题1：（1）BC=CF+CD；（2）BC=CF-CD，各6分。题2：（1）∠BME=∠CNE，6分；（2）OM=ON，6分；（3）△AGD直角，直角在G，12分。第三问建议：构造T及中点4分、△TAB等边2分、△AGF等边2分、等腰角度与直角结论4分。'),
    ('来源与文件', '原题、原图与固定版本出处保留在README及sources。本版图形为独立教学重绘。除这份完整教师版，另提供只含三问压轴题的challenge_detailed.pdf、Word和在线图文版，便于手机阅读。'),
]


def word_and_markdown(b,challenge_only):
    name='challenge_detailed' if challenge_only else 'extra_teacher'
    title='三问压轴题：一页一步教师详解' if challenge_only else '两道补充题：逐步配图教师详解'
    doc=base.word_base(title)
    doc.sections[0].footer.paragraphs[0].text='详解修订版｜2026-10-17｜出处与编排说明见README'
    normal=doc.styles['Normal'];normal.paragraph_format.space_after=Pt(4)
    normal.paragraph_format.line_spacing=1.12
    parts=['# '+title,'题目来源与编排说明见[README](README.md)。']
    doc.add_heading(title,0)
    if not challenge_only:
        for title,body in INTRO:
            doc.add_heading(title,2);doc.add_paragraph(body);parts.extend(['## '+title,body])
        for s in b.MEDIUM_STEPS:
            doc.add_page_break();doc.add_heading('题1｜'+s['title'],1)
            doc.add_paragraph(s['action']);doc.add_picture(str(b.ROOT/'teaching_figures'/f"{s['key']}.png"),width=Cm(10.5))
            doc.add_paragraph(s['proof']);doc.add_paragraph('结论：'+s['result'])
            parts.extend(['## 题1｜'+s['title'],f"![步骤图](teaching_figures/{s['key']}.png)",s['proof'],'**结论：** '+s['result']])
        doc.add_page_break()
    doc.add_heading('三问压轴题｜题面与读图',1);doc.add_paragraph(b.HARD)
    doc.add_picture(str(b.ROOT/'sources/challenge_630.png'),width=Cm(16))
    parts.extend(['## 三问压轴题｜完整题面',b.HARD,'![原题三图](sources/challenge_630.png)',
                  '阅读顺序：未学中点连线性质先看两页引理；已学可直接从第（1）问开始。原题保持不变，下面重写完整解释。'])
    for s in LESSONS:
        doc.add_page_break();doc.add_heading(s['group']+'｜'+s['title'],1)
        doc.add_paragraph('为什么这样想：'+s['why']);doc.add_picture(str(b.ROOT/'teaching_figures'/f"{s['key']}.png"),width=Cm(10.8))
        doc.add_paragraph('看图／作图：'+s['action'])
        doc.add_paragraph(s['proof']);doc.add_paragraph('本页结论：'+s['result']);doc.add_paragraph('讲解提醒：'+s['check'])
        parts.extend(['## '+s['group']+'｜'+s['title'],'**为什么这样想：** '+s['why'],
                      f"![本步放大图](teaching_figures/{s['key']}.png)",'**看图／作图：** '+s['action'],
                      s['proof'].replace('\n','\n\n'),'**本页结论：** '+s['result'],'**讲解提醒：** '+s['check']])
    doc.add_page_break();doc.add_heading('三问复盘与教师追问',1)
    parts.append('## 三问复盘与教师追问')
    for title,body in RECAP:
        doc.add_heading(title,2);doc.add_paragraph(body);parts.extend(['### '+title,body.replace('\n','\n\n')])
    doc.save(b.ROOT/(name+'.docx'))
    (b.ROOT/(name+'.md')).write_text('\n\n'.join(parts)+'\n',encoding='utf-8')


def build(b):
    render_figures(b)
    for challenge_only in [False,True]:
        path=b.ROOT/('challenge_detailed.pdf' if challenge_only else 'extra_teacher.pdf')
        c=canvas.Canvas(str(path),pagesize=b.PAGE);page=1
        if not challenge_only:
            heading(c,'教师版｜本次修订与课堂安排',page);y=252*MM
            for title,body in INTRO:
                y=base.txt(c,title,18*MM,y,size=11.5,lead=18)
                y=base.txt(c,body,18*MM,y-5,size=10.5,lead=17)-12
            if y<24*MM:raise ValueError('Intro overflow')
            c.showPage();page+=1
            heading(c,'题1：中等题｜原解保留',page)
            b.block(c,b.MEDIUM_STEPS[0],250);b.block(c,b.MEDIUM_STEPS[1],140)
            c.showPage();page+=1
        overview(c,page,b);page+=1
        for s in LESSONS:
            lesson_page(c,s,page,b);page+=1
        recap(c,page);c.save()
        word_and_markdown(b,challenge_only)
