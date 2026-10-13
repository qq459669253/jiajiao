"""Selected published problems, not newly invented questions (2026-10-11).

Attribution, upstream revisions, modifications and component licenses: README.md.
Run from any directory: python build_selected_pack.py
Dependencies: reportlab, python-docx, pymupdf.
"""
from pathlib import Path
import zipfile
import argparse
import fitz
import reportlab
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
FIG = ROOT / 'sources' / 'figures'
pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
pdfmetrics.registerFont(TTFont('MathGlyphs', str(Path(reportlab.__file__).parent/'fonts'/'Vera.ttf')))
FONT = 'STSong-Light'
MM = 72 / 25.4
W, H = 210 * MM, 297 * MM

PROBLEMS = [
    {
        'title': '1．旋转与三段距离和的最小值（8分，建议10分钟）',
        'source': '《初中数学总复习讲义》1.5.1，例6.17.2，印刷页57（PDF第62页）。',
        'text': '如图，正方形ABCD以AB为边向外作等边三角形ABE，M为对角线BD上一点，将BM绕B逆时针旋转60°到BN，若正方形边长为2，求AM+BM+CM的最小值。',
        'figures': ['onion_612.png'],
        'note': '作答要求：写出长度如何转移、最小值及其取到的位置。若未学勾股定理，可用图中的一条固定线段表示最小值，根式计算不计入本次60分。',
    },
    {
        'title': '2．等边三角形动点与中点轨迹（10分，建议10分钟）',
        'source': '《初中数学总复习讲义》1.5.1，习题6.25.12，印刷页79（PDF第84页）。',
        'text': '如图，C、D为线段AB上两点，P是线段CD上的动点，分别以AP、BP为边在AB同侧作两个等边三角形APE和BPF，M是EF的中点。已知AB=20，AC=BD=2，当P从C运动到D时（无重复运动），求M点的运动路径长。',
        'figures': ['onion_632.png'],
        'note': '作答要求：说明M沿什么路径运动，并证明路径长。不能只按图形比例猜测。',
    },
    {
        'title': '3．半角模型：从正方形到一般图形（20分，建议15分钟）',
        'source': '《初中数学总复习讲义》1.5.1，习题6.16.3，印刷页56（PDF第61页）。',
        'text': '（1）如图1，在正方形ABCD中，E、F为BC、CD上点，且∠EAF=45°，直接写出BE、DF、EF三条线段的数量关系。\n（2）如图2，将正方形ABCD改为四边形ABCD，AB=AD，∠B+∠D=180°，E、F分别为边BC、CD上的点，∠EAF=(1/2)∠BAD，判断（1）中结论是否依然成立。\n（3）在（2）基础上将△AEF绕A逆时针旋转，E、F运动到BC、CD延长线上，如图3，判断BE、DF、EF三条线段的数量关系并证明。',
        'figures': ['onion_6107.jpg'],
        'note': '本次评分：（1）4分；（2）8分，除判断外写出理由；（3）8分。图3中B、C、E依次共线，C、D、F依次共线。',
    },
    {
        'title': '4．直角旋转、中点与面积最值（22分，建议15分钟）',
        'source': '公开题库middleschoolmath-ebook所标“2017年河南中考”，题库编号T22-A01-10。',
        'text': '如图1，在Rt△ABC中，∠A=90°，AB=AC，点D、E分别在边AB、AC上，AD=AE，连接DC，点M、P、N分别为DE、DC、BC的中点。\n（1）观察猜想：图1中，线段PM与PN的数量关系是__________，位置关系是__________。\n（2）探究证明：把△ADE绕点A逆时针方向旋转到图2的位置，连接MN、BD、CE，判断△PMN的形状，并说明理由。\n（3）拓展延伸：把△ADE绕点A在平面内自由旋转，若AD=4，AB=10，请直接写出△PMN面积的最大值。',
        'figures': ['henan2017_01.png', 'henan2017_02.png'],
        'note': '本次评分：（1）4分；（2）12分；（3）6分。第（3）问请补充为什么达到最大值及此时D的位置，便于诊断推理过程。',
    },
]

GUIDE = [
    ('如何用这套现有题做两小时诊断', [
        ('选题目的', '针对原卷选择第10题、填空第16题及第24、25题的失分，主要检测旋转后的长度转移、中点辅助线、位置变化时的分类，以及最值的取等条件。题目选自现成资料，保留原条件、数值和原图；60分评分、时间、作答要求及以下讲解为本次编排。'),
        ('难度与进度', '这是一套提高诊断卷，不是基础题检测。第1题需要把三段长度转成一条折线；第2题需要主动补一个固定等边三角形；第3题需两次全等并在延长线位置重新判断加减；第4题需连接中点性质、旋转和面积。不要把“中考题”标签当作难度保证，重在是否独立完成这些方法链。'),
        ('先备知识', '只用全等、角关系、等边三角形、平行线、三角形不等式与面积公式完成主要推理。正方形只需四边相等、四角直角。中点连线性质若尚未学习，可用第6页的两次SAS证明；记录为先备知识缺口，不直接判断为思维薄弱。第1题的根式数值另需勾股定理，不计本次评分。'),
        ('120分钟安排', '0—5分钟交代规则并确认先备知识；5—55分钟独立完成主卷；55—65分钟批改并追问卡点；65—105分钟讲解两条主要方法链；105—117分钟遮住原解，独立重建关键证明；117—120分钟让孩子复述方法并记录下一课目标。'),
        ('讲解取舍', '若第1、2题较弱：18分钟讲“旋转搬长度+等号条件”，18分钟讲“固定图形+中点轨迹”，4分钟复述。若小题完成较好：18分钟讲第3题两次全等及延长线加减，18分钟讲第4题旋转+中点+面积，4分钟复述。若中点性质不会，先用8分钟补引理，余下32分钟只精讲一个核心大题。'),
        ('使用边界', '原卷已被学校讲评，做过同型题不等于掌握。这次优先观察辅助线是否自己想到、夹角能否说明，以及最值能否证明取得。50分钟结束可保留未完成，不用把60分线性换算成150分预测成绩。'),
    ]),
    ('第1题：三段距离变成一条折线', [
        ('答案', '最小值为CE，数值为√6+√2。M取直线EC与对角线BD的交点；这是BD内部的点。主要评分只要求以CE表示，并证明下界能够取得。'),
        ('把AM搬到EN', 'BA=BE，BM=BN，∠ABM=∠EBN=45°，所以△ABM≌△EBN（SAS），得AM=EN。由BM=BN、∠MBN=60°，△BMN为等边三角形，故BM=MN。当M=B时直接处理，三段和为4，大于最优值。'),
        ('固定下界与取等', '当M≠B时，AM+BM+CM=EN+NM+MC≥EC。等号要求E、N、M、C在同一直线上，且按此顺序排列，不能只写一个三角形不等式就结束。'),
        ('为什么这样的M合法', '在等腰△EBC中，BE=BC=2，∠EBC=150°，所以∠ECB=15°。又BD与BC夹45°，而等边△BMN使MN与BC夹15°，方向与EC相同，故MN∥EC。令M为EC与BD的交点，则MN所在直线就是EC。射线BE、BN、BM、BC在∠EBC内部依次排列，它们与线段EC的交点依次为E、N、M、C。EC与BD交在正方形内部，因而M在BD上，等号可取得。'),
        ('数值计算，按进度选讲', '作EH⊥AB于H，H为AB中点，因此BH=1、EH=√3（这里使用勾股定理）。相对BC方向，E的垂直距离为1，水平方向到C的距离为2+√3，所以CE²=(2+√3)²+1=8+4√3，CE=√6+√2。未学勾股的学生停在“最小值为CE”即可。'),
        ('8分评分与提示', 'AM=EN的全等2分；BM=MN及折线转化2分；固定下界2分；合法取点、顺序及等号2分。提示顺序：①哪条长度能旋转到N附近？②BMN是什么三角形？③折线的端点是否固定？④等号时M落在哪里？记录提示等级。'),
    ]),
    ('第2题：先补固定等边三角形', [
        ('答案', '运动路径是一条与AB平行的线段，长为8。只答8不能确认会做动点迁移。'),
        ('找到固定点Q', '在AB同侧作固定的等边△ABQ。由等边三角形的60°角可知E在AQ上，F在BQ上，EP∥BQ。又EP=AP，FQ=BQ-BF=AB-BP=AP，所以EP=FQ。此处没有用相似比例。'),
        ('证明M是QP中点', 'M为EF中点，EM=FM；EP∥FQ使∠PEM=∠QFM；EP=FQ，所以△PEM≌△QFM（SAS），得PM=QM。对应角又说明PM、QM为反向共线的两条射线，故Q、M、P共线，M为QP中点。这里要同时证明共线和等长。'),
        ('再判断轨迹及长度', '设H为QA中点。对任意P，△QAP中H、M为两边中点，所以HM∥AP，即M在过H且平行AB的固定直线上。P从C到D时，M分别从QC的中点K到QD的中点L。△QCD的中点连线给KL=CD/2；CD=20-2-2=16，所以KL=8。对应P的顺序保持，M沿KL单向运动，无额外折返。'),
        ('10分评分与提示', '补固定Q并找出EP=FQ、平行关系2分；SAS且说明M为QP中点4分；固定直线、端点和半长关系2分；答案8得2分。提示顺序：①动三角形旁边能否补一个固定三角形？②比较EP和FQ。③已知EF中点，哪两三角形能全等？'),
        ('诊断区分', '若给出Q后能独立证明，主要卡在辅助线选择；若已指出M是QP中点仍无法描述轨迹，主要卡在固定点和动点关系；若中点连线性质未学，先补引理，再观察能否自行接上。'),
    ]),
    ('第3题：同一辅助线，关系由加变减', [
        ('答案', '（1）EF=BE+DF。（2）仍成立。（3）BE=DF+EF，亦即EF=BE-DF。第（3）问不能照搬前两问的加法关系。'),
        ('第（1）（2）问：第一次全等', '沿CB过B延长到G，使BG=DF。由∠ABG=180°-∠ABC=∠ADC=∠ADF，AB=AD，BG=DF，得△ABG≌△ADF（SAS）。于是AG=AF，∠BAG=∠DAF。正方形是这一论证的特殊情形。'),
        ('第二次全等与加法', '图1、2中，∠GAE=∠GAB+∠BAE=∠DAF+∠BAE=∠BAD-∠EAF=∠EAF。又AG=AF，AE公共，故△GAE≌△FAE（SAS），得到GE=FE。此时G、B、E依次共线，所以GE=GB+BE=DF+BE。'),
        ('第（3）问：改变复制长度的位置', '这次在射线BC上取G，使BG=DF。F在CD越过D的延长线上，∠ADF=180°-∠ADC=∠ABC=∠ABG。由AB=AD、BG=DF，仍有△ABG≌△ADF（SAS），AG=AF，∠BAG=∠DAF。'),
        ('第二次全等与减法', '图3中，∠BAE=∠EAF+∠DAF。因此∠GAE=∠BAE-∠BAG=∠EAF，结合AG=AF、AE公共，得△GAE≌△FAE（SAS），GE=EF。因为0<∠BAG=∠DAF<∠BAE，射线AG在∠BAE内部，它与BE相交于G，故B、G、E依次共线。于是GE=BE-BG=BE-DF，得到BE=DF+EF。'),
        ('20分评分与提示', '（1）正确关系4分。（2）判断2分，复制长度与第一次全等3分，第二次全等及结论3分。（3）关系2分，复制长度及第一次全等2分，第二次全等2分，点的顺序与减法2分。提示先问“需要搬哪一段”，再问“复制长度放在B的哪边”，最后才给G。'),
        ('对应原卷的能力', '这题保留了“基本位置—条件推广—图形换位置”的层次，接近原卷第24、25题的后续小问。关键是每换一次位置，都重新检查夹角和线段顺序。不要只背“半角模型等于和”。'),
    ]),
    ('第4题：旋转等长垂直，再连接中点', [
        ('答案', '（1）PM=PN，PM⊥PN。（2）△PMN为等腰直角三角形，直角顶点为P。（3）最大面积49/2，D在射线AB的反向延长线上，即D、A、B依次共线时取得。'),
        ('长度与方向先来自哪里', '将平面绕A逆时针旋转90°，B恰好转到C，D恰好转到E，因此线段BD转到CE。得到BD=CE且BD⊥CE。对图2也可先说明∠BAD=∠CAE，AB=AC、AD=AE，从△BAD≌△CAE（SAS）得BD=CE；90°旋转同时确定对应线段的方向。两组三角形必须同向旋转，不能只看到两个直角就猜结论。'),
        ('把结论搬到PM、PN', '在△DCE中，P、M分别为DC、DE中点，所以PM∥CE，PM=CE/2；在△DCB中，P、N分别为DC、CB中点，所以PN∥DB，PN=DB/2。结合BD=CE、BD⊥CE，得到PM=PN、PM⊥PN。所以△PMN为等腰直角三角形。这些中点结论若未学，先看下一页证明。'),
        ('面积为何有固定上界', '自由旋转时，等长和垂直关系仍成立，故S△PMN=(1/2)PM·PN=BD²/8。三角形不等式给BD≤AB+AD=14，所以面积≤14²/8=49/2。等号在D、A、B依次共线时成立；自由旋转允许这一位置，所以是可取得的最大值。不能只列出14而不说明旋转能否到达。'),
        ('退化位置须说明', 'BD=0不可能，因为AB=10、AD=4。个别旋转位置可能使C、D、E或C、D、B共线，但其中点平行和半长关系直接由共线线段成立，或由旋转连续性得到；不影响该题最大值。计算最大值的D、A、B共线位置下△PMN并不退化。'),
        ('22分评分与提示', '（1）两关系各2分。（2）BD=CE的全等或明确旋转理由4分；BD⊥CE的方向理由2分；两条中点连线4分；等腰直角结论2分。（3）面积化为BD²/8得2分；上界及49/2得2分；合法取等位置2分。'),
        ('防止机械套用', '只写“中点所以一半”而说不出是哪两个三角形，说明结构识别不足；只证明等长但没证明垂直，说明还没把旋转方向与对应角结合；只写49/2而不知道D的位置，说明最值等号仍需补课。'),
    ]),
    ('中点引理、课堂回练与出处核对', [
        ('中点连线引理：用两次SAS补齐', '△TUV中，X、Y分别为TU、TV中点。延长XY至Z，使YZ=YX。由TY=YV、YX=YZ及对顶角，△TYX≌△VYZ（SAS），故VZ=TX=UX，且VZ∥UX。再由UX=VZ、UZ公共、∠XUZ=∠VZU，△XUZ≌△VZU（SAS），所以XZ=VU，且XZ∥VU。因XZ=2XY，得XY∥UV、XY=UV/2。先让孩子自己写两组全等式，再应用于第2、4题。'),
        ('12分钟回练：不再另编新题', '遮住答案，按讲解内容选两项：A．第1题只给“AM=EN”，重新写出剩下的下界、取等顺序及M的位置；B．第2题只给固定Q，重新证明M是QP中点并求轨迹长；C．第3题只看图3，独立决定G放在哪一侧并写出两次SAS；D．第4题只看第（3）问，独立写出面积式、上界和等号位置。约8分钟重做、4分钟核对。这里是原题关键步骤的回练，不能将短时间重做成功视为陌生题迁移已完全掌握。'),
        ('课后再测', '下一次应选一份未见过的同型现成题再测，以确认迁移。本包不把改数字的自编题冒充公开题，也不将同题即时重做当作第二套独立测试。'),
        ('出处与核对程度', '第1—3题来自Onion12138/math的《初中数学总复习讲义》1.5.1，有原图和可定位的例题、习题号；第4题来自higyg/middleschoolmath-ebook的T22-A01-10，原题库明确标注2017年河南中考，原解给出了PM=PN、PM⊥PN及49/2。本包独立补全证明，但未取得官方试卷对照，不宣称学校来源或官方授权。README提供固定版本原链接、题目摘录和许可文件。'),
        ('记录表', '独立摸底____/60；自想到的辅助线：________。\n卡点：□夹角 □复制长度 □旋转对应 □中点引理 □线段顺序 □等号位置。\n提示等级：□独立 □口头方向 □给出辅助线 □完整示范。\n回练完成：________；下一课要用陌生题再测：________。'),
    ]),
]


def normalized(s):
    return s.replace('−', '-').replace('↔', '对应')


def glyph_font(ch):
    return 'MathGlyphs' if ch in '√²' else FONT


def measured(s, size):
    return sum(pdfmetrics.stringWidth(ch, glyph_font(ch), size) for ch in s)


def wrap(s, width, size):
    result = []
    for para in normalized(s).split('\n'):
        buf = ''
        for ch in para:
            if buf and measured(buf + ch, size) > width:
                result.append(buf)
                buf = ch
            else:
                buf += ch
        result.append(buf)
    return result


def txt(c, s, x, y, width=174*MM, size=10.5, lead=17):
    c.setFillColorRGB(.06, .09, .12)
    for line in wrap(s, width, size):
        xx = x
        current_font = None
        for ch in line:
            font = glyph_font(ch)
            if font != current_font:
                c.setFont(font, size)
                current_font = font
            c.drawString(xx, y, ch)
            xx += pdfmetrics.stringWidth(ch, font, size)
        y -= lead
    return y


def header(c, title, subtitle, page, revision='2026-10-11'):
    txt(c, title, 18*MM, 279*MM, size=17, lead=20)
    txt(c, subtitle, 18*MM, 270*MM, size=9, lead=13)
    c.setLineWidth(.4)
    c.setStrokeColorRGB(.4, .45, .5)
    c.line(18*MM, 264*MM, 192*MM, 264*MM)
    txt(c, '现有题选编｜出处与编排说明见同目录README｜'+revision, 18*MM, 12*MM, size=8, lead=10)
    c.setFont(FONT, 9)
    c.drawRightString(192*MM, 12*MM, str(page))


def pic(c, file, x, y, w, h):
    img = ImageReader(str(FIG / file))
    iw, ih = img.getSize()
    scale = min(w / iw, h / ih)
    dw, dh = iw * scale, ih * scale
    c.drawImage(img, x + (w-dw)/2, y + (h-dh)/2, dw, dh, mask='auto')


def blanks(c, y, n, width=174*MM):
    c.setLineWidth(.3)
    c.setStrokeColorRGB(.76, .79, .82)
    for i in range(n):
        yy = y - 9*MM*i
        if yy < 23*MM:
            raise ValueError('Answer lines cross footer')
        c.line(18*MM, yy, 18*MM+width, yy)


def render_source_figures():
    for suffix in ['01', '02']:
        doc = fitz.open(FIG / f'henan2017_{suffix}.pdf')
        doc[0].get_pixmap(matrix=fitz.Matrix(3, 3)).save(FIG / f'henan2017_{suffix}.png')


def student_pdf():
    c = canvas.Canvas(str(ROOT/'selected_student.pdf'), pagesize=(W, H))
    c.setTitle('初二全等与旋转进阶摸底：现有题选编（学生版）')
    for i, q in enumerate(PROBLEMS, 1):
        header(c, '初二全等与旋转：现有题进阶摸底', '姓名：________　独立限时50分钟　共60分　请保留思路和辅助线', i)
        y = txt(c, q['title'], 18*MM, 255*MM, size=12, lead=20)
        y = txt(c, q['source'], 18*MM, y-3, size=8.5, lead=13)
        y = txt(c, q['text'], 18*MM, y-8, size=11, lead=18)
        if i == 1:
            pic(c, q['figures'][0], 106*MM, 166*MM, 83*MM, 51*MM)
            txt(c, q['note'], 18*MM, 204*MM, width=83*MM, size=9.5, lead=16)
            blanks(c, 158*MM, 14)
        elif i == 2:
            pic(c, q['figures'][0], 109*MM, 163*MM, 80*MM, 53*MM)
            txt(c, q['note'], 18*MM, 204*MM, width=83*MM, size=9.5, lead=16)
            blanks(c, 155*MM, 14)
        elif i == 3:
            pic(c, q['figures'][0], 18*MM, 134*MM, 174*MM, 54*MM)
            txt(c, q['note'], 18*MM, 125*MM, size=9, lead=15)
            blanks(c, 109*MM, 10)
        else:
            pic(c, q['figures'][0], 19*MM, 113*MM, 78*MM, 55*MM)
            pic(c, q['figures'][1], 108*MM, 113*MM, 78*MM, 55*MM)
            txt(c, q['note'], 18*MM, 105*MM, size=9, lead=15)
            blanks(c, 87*MM, 8)
        c.showPage()
    header(c, '证明续页：现有题进阶摸底', '可用于第3、4题；请标清题号并画出所用辅助线', 5)
    blanks(c, 254*MM, 25)
    c.save()


def teacher_pdf():
    import illustrated_teacher as illustrated
    illustrated.figures(ROOT, txt)
    illustrated.pdf(ROOT, GUIDE, txt, header)
    illustrated.pdf(ROOT, GUIDE, txt, header, only_q4=True)


def word_base(title):
    d = Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    sec.left_margin = sec.right_margin = Cm(1.8)
    style = d.styles['Normal']
    style.font.name = '宋体'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(6)
    sec.footer.paragraphs[0].text = '现有题选编｜出处、改动及许可见README｜2026-10-11'
    d.core_properties.title = title
    d.core_properties.subject = '出处与编排说明见同目录README.md；题目条件与数值保留。'
    return d


def word_files():
    d = word_base('初二全等与旋转进阶摸底：现有题选编')
    for i, q in enumerate(PROBLEMS, 1):
        if i > 1:
            d.add_page_break()
        d.add_heading('初二全等与旋转：现有题进阶摸底', level=0)
        d.add_paragraph('姓名：________　限时50分钟　共60分　请保留辅助线和证明')
        d.add_heading(q['title'], level=1)
        d.add_paragraph(q['source'])
        d.add_paragraph(q['text'])
        for file in q['figures']:
            width = 15.5 if i == 3 else (7.5 if i == 4 else 8)
            d.add_picture(str(FIG/file), width=Cm(width))
        d.add_paragraph(q['note'])
        for _ in range(5):
            d.add_paragraph('________________________________________________________________')
    d.add_page_break()
    d.add_heading('证明续页（请标明题号）', level=1)
    for _ in range(20):
        d.add_paragraph('________________________________________________________________')
    d.save(ROOT/'selected_student.docx')
    import illustrated_teacher as illustrated
    illustrated.word(ROOT, GUIDE, word_base)
    illustrated.word(ROOT, GUIDE, word_base, only_q4=True)


def markdown_files():
    import illustrated_teacher as illustrated
    illustrated.markdown(ROOT, GUIDE)
    illustrated.markdown(ROOT, GUIDE, only_q4=True)
    text = '# 学生题面文本源\n\n这是现有题的中文重排文本。具体出处及改动见README.md，原始TeX摘录及原图见sources。\n\n'
    for q in PROBLEMS:
        text += f"## {q['title']}\n\n{q['source']}\n\n{q['text']}\n\n{q['note']}\n\n"
    (ROOT/'selected_student.md').write_text(text.rstrip()+'\n', encoding='utf-8')


def bundle():
    archive = ROOT/'selected_existing_problems.zip'
    with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED) as z:
        for path in sorted(ROOT.rglob('*')):
            if path.is_file() and path != archive and '__pycache__' not in path.parts:
                z.write(path, path.relative_to(ROOT))


if __name__ == '__main__':
    parser=argparse.ArgumentParser(description='Build selected math assessment materials.')
    parser.add_argument('--teacher-only',action='store_true',help='Update illustrated teacher files, preserving student documents.')
    args=parser.parse_args()
    if args.teacher_only:
        import illustrated_teacher as illustrated
        teacher_pdf()
        illustrated.word(ROOT, GUIDE, word_base)
        illustrated.word(ROOT, GUIDE, word_base, only_q4=True)
        illustrated.markdown(ROOT, GUIDE)
        illustrated.markdown(ROOT, GUIDE, only_q4=True)
    else:
        render_source_figures()
        student_pdf()
        teacher_pdf()
        word_files()
        markdown_files()
    bundle()
    print('Built student PDF/Word, teacher PDF/Word/Markdown and ZIP.')
