"""GPL-3.0: step-by-step teaching diagrams and solutions, 2026-10-13.

These are newly drawn mathematical explanation figures. Original question
conditions and original student figures remain in build_selected_pack.py.
"""
from pathlib import Path
import math
import io
import fitz
from reportlab.pdfgen import canvas
from docx.shared import Cm, Pt

MM = 72/25.4
PAGE = (210*MM, 297*MM)
RED = (.72, .12, .14)
BLUE = (.05, .32, .66)
GREEN = (.08, .43, .31)
GRAY = (.52, .55, .58)
INK = (.12, .15, .18)


def add(a, b): return (a[0]+b[0], a[1]+b[1])
def sub(a, b): return (a[0]-b[0], a[1]-b[1])
def mul(a, t): return (a[0]*t, a[1]*t)
def mid(a, b): return mul(add(a, b), .5)
def length(a): return math.hypot(*a)
def cross(a, b): return a[0]*b[1]-a[1]*b[0]
def rotate(a, t): return (a[0]*math.cos(t)-a[1]*math.sin(t), a[0]*math.sin(t)+a[1]*math.cos(t))
def meet(a, b, c, d):
    v, w = sub(b, a), sub(d, c)
    return add(a, mul(v, cross(sub(c, a), w)/cross(v, w)))


def base_q4(rotated=True):
    p = {'A': (0., 4.), 'B': (-4., 0.), 'C': (4., 0.)}
    p['D'] = (-.25, 2.) if rotated else (-1.8, 2.2)
    p['E'] = add(p['A'], rotate(sub(p['D'], p['A']), math.pi/2))
    p['M'], p['P'], p['N'] = mid(p['D'], p['E']), mid(p['D'], p['C']), mid(p['B'], p['C'])
    return p


def base_q3(kind):
    if kind == 'square':
        p = {'A': (0., 2.), 'B': (0., 0.), 'C': (2., 0.), 'D': (2., 2.)}
        p['E'] = (.8, 0.)
        ang = math.pi/4
    else:
        p = {'A': (0., 0.), 'B': (0., -2.), 'D': (2*math.cos(.35), 2*math.sin(.35))}
        ox = (2+p['D'][1])/p['D'][0]
        radius = math.hypot(ox, -1)
        p['C'] = (ox+radius*math.cos(5.25), -1+radius*math.sin(5.25))
        t = .4 if kind == 'general' else 3.5
        p['E'] = add(p['B'], mul(sub(p['C'], p['B']), t))
        ang = (math.pi/2+.35)/2
    direction = rotate(sub(p['E'], p['A']), ang)
    p['F'] = meet(p['A'], add(p['A'], direction), p['C'], p['D'])
    df = length(sub(p['F'], p['D']))
    bc = sub(p['C'], p['B'])
    sign = 1 if kind == 'outside' else -1
    p['G'] = add(p['B'], mul(bc, sign*df/length(bc)))
    return p


class Drawing:
    """Vector geometry with equal scales on both axes and explicit marks."""
    def __init__(self, c, pts, txt):
        self.c, self.txt = c, txt
        xs, ys = [p[0] for p in pts.values()], [p[1] for p in pts.values()]
        scale = min(214/(max(xs)-min(xs)), 166/(max(ys)-min(ys)))
        ox = (260-scale*(max(xs)-min(xs)))/2
        oy = (220-scale*(max(ys)-min(ys)))/2
        self.p = {k: (ox+(p[0]-min(xs))*scale, oy+(p[1]-min(ys))*scale) for k,p in pts.items()}

    def segment(self, a, b, color=GRAY, width=.85, dash=False):
        self.c.saveState()
        self.c.setStrokeColorRGB(*color)
        self.c.setLineWidth(width)
        if dash: self.c.setDash(4, 3)
        self.c.line(*self.p[a], *self.p[b])
        self.c.restoreState()

    def edges(self, pairs, **kw):
        for a, b in pairs: self.segment(a, b, **kw)

    def triangle(self, names, color):
        self.c.saveState()
        self.c.setFillColorRGB(*color)
        self.c.setFillAlpha(.10)
        path = self.c.beginPath()
        path.moveTo(*self.p[names[0]])
        for k in names[1:]: path.lineTo(*self.p[k])
        path.close()
        self.c.drawPath(path, fill=1, stroke=0)
        self.c.restoreState()
        self.edges(list(zip(names, names[1:]+names[:1])), color=color, width=1.45)

    def tick(self, a, b, n=1, color=INK):
        pa, pb = self.p[a], self.p[b]
        v = sub(pb, pa); v = mul(v, 1/length(v)); normal = (-v[1], v[0])
        center = mid(pa, pb)
        for i in range(n):
            t = (i-(n-1)/2)*3
            center2 = add(center, mul(v, t))
            self.c.setStrokeColorRGB(*color)
            self.c.setLineWidth(1)
            self.c.line(*add(center2, mul(normal, -3.5)), *add(center2, mul(normal, 3.5)))

    def right(self, v, a, b, color=INK):
        p, u, w = self.p[v], sub(self.p[a], self.p[v]), sub(self.p[b], self.p[v])
        u, w = mul(u, 7/length(u)), mul(w, 7/length(w))
        self.c.setStrokeColorRGB(*color)
        self.c.setLineWidth(.9)
        self.c.line(*add(p, u), *add(add(p, u), w))
        self.c.line(*add(p, w), *add(add(p, u), w))

    def angle(self, v, a, b, label, color=BLUE, radius=18):
        p = self.p[v]
        u, w = sub(self.p[a], p), sub(self.p[b], p)
        start, end = math.degrees(math.atan2(u[1], u[0])), math.degrees(math.atan2(w[1], w[0]))
        delta = (end-start+180)%360-180
        self.c.setStrokeColorRGB(*color)
        self.c.setLineWidth(.9)
        self.c.arc(p[0]-radius, p[1]-radius, p[0]+radius, p[1]+radius, start, delta)
        theta = math.radians(start+delta/2)
        pos = add(p, ((radius+11)*math.cos(theta), (radius+11)*math.sin(theta)))
        self.txt(self.c, label, pos[0]-len(label)*2.4, pos[1]-3, width=80, size=9, lead=11)

    def label_line(self, a, b, s, offset=(0, 8), color=BLUE):
        pos = add(mid(self.p[a], self.p[b]), offset)
        self.txt(self.c, s, pos[0]-len(s)*2.5, pos[1], width=100, size=10, lead=12)

    def labels(self, offsets=None):
        offsets = offsets or {}
        self.c.setFont('Helvetica', 11)
        self.c.setFillColorRGB(*INK)
        for k, pos in self.p.items():
            dx, dy = offsets.get(k, (-5, 7))
            self.c.circle(*pos, 1.2, fill=1, stroke=0)
            self.c.drawString(pos[0]+dx, pos[1]+dy, k)


def draw_q4(c, key, txt):
    initial = key in ['q4a1', 'q4a2', 'q4a3']
    p = base_q4(not initial)
    if key == 'q4max':
        p = {'A': (0.,0.), 'B': (10.,0.), 'C': (0.,10.), 'D': (-4.,0.), 'E': (0.,-4.)}
        p.update(M=mid(p['D'],p['E']), P=mid(p['D'],p['C']), N=mid(p['B'],p['C']))
    if key == 'q4angles': p['T'] = meet(p['B'],p['D'],p['C'],p['E'])
    if key in ['q4sas','q4angles']:
        p={k:v for k,v in p.items() if k not in ['M','P','N']}
    elif key in ['q4a1','q4pm']:
        p.pop('N')
    elif key in ['q4a2','q4pn']:
        p.pop('M')
    if key == 'q4bound':
        p = {k:p[k] for k in ['A','B','D']}
        d = Drawing(c,p,txt)
        d.triangle(['B','A','D'],BLUE)
        d.label_line('A','B','10',(-9,6)); d.label_line('A','D','4',(12,0))
        d.label_line('B','D','BD≤14',(0,-13))
        d.labels({'A':(-4,9),'B':(-16,-6),'D':(7,-5)})
        return
    d = Drawing(c,p,txt)
    d.edges([('A','B'),('B','C'),('C','A'),('A','D'),('D','E'),('E','A'),('D','C')])
    offsets = {'A':(-5,9),'B':(-14,-12),'C':(5,-12),'D':(-16,-3),'E':(5,3),'M':(-13,5),'P':(5,-3),'N':(-3,-14),'T':(6,-1)}
    if key in ['q4a1','q4pm']:
        d.triangle(['D','C','E'],BLUE)
        d.segment('P','M',RED,2)
        for a,b in [('D','P'),('P','C')]:d.tick(a,b,1)
        for a,b in [('D','M'),('M','E')]:d.tick(a,b,2)
    elif key in ['q4a2','q4pn']:
        d.triangle(['D','C','B'],BLUE)
        d.segment('P','N',RED,2)
        for a,b in [('D','P'),('P','C')]:d.tick(a,b,1)
        for a,b in [('B','N'),('N','C')]:d.tick(a,b,2)
    elif key == 'q4sas':
        d.triangle(['B','A','D'],RED);d.triangle(['C','A','E'],BLUE)
        d.tick('A','B',1);d.tick('A','C',1);d.tick('A','D',2);d.tick('A','E',2)
        d.angle('A','B','D','',RED,20);d.angle('A','C','E','',BLUE,30)
    elif key == 'q4angles':
        d.triangle(['B','C','T'],GREEN)
        d.segment('B','D',RED,1.6);d.segment('D','T',RED,1.6,True)
        d.segment('C','E',BLUE,1.8);d.right('T','B','C',GREEN)
        d.angle('B','A','D','x',RED,13);d.angle('C','A','E','x',BLUE,13)
        d.angle('B','C','T','45°-x',GREEN,30);d.angle('C','B','T','45°+x',GREEN,30)
        offsets['P']=(5,-12);offsets['E']=(5,10)
    else:
        d.segment('B','D',BLUE,1.3);d.segment('C','E',BLUE,1.3)
        d.triangle(['M','P','N'],RED)
        d.right('P','M','N',RED);d.tick('P','M',1,RED);d.tick('P','N',1,RED)
        if key == 'q4area':
            d.label_line('P','M','BD/2',(-22,0));d.label_line('P','N','BD/2',(12,-14))
        if key == 'q4max':
            d.segment('D','B',BLUE,1.6);d.segment('E','C',BLUE,1.6)
            d.label_line('D','A','4',(0,-13));d.label_line('A','B','10',(0,-13))
            d.label_line('P','M','7',(-12,0));d.label_line('P','N','7',(0,8))
            offsets.update(A=(5,-13), B=(5,-3), C=(-10,8), D=(-14,-3), E=(4,-14),M=(-17,-10),P=(-17,3),N=(5,5))
    d.labels(offsets)


def draw_q1(c,key,txt):
    s=math.sqrt(3)
    p={'A':(0.,2.),'B':(0.,0.),'C':(2.,0.),'D':(2.,2.),'E':(-s,1.)}
    t=1-1/s if key=='q1equal' else .8
    p['M']=(t,t);p['N']=rotate(p['M'],math.pi/3)
    if key=='q1number':
        p={k:v for k,v in p.items() if k not in ['M','N']}
        p.update(H=(0.,1.),K=(-s,0.))
    d=Drawing(c,p,txt)
    d.edges([('A','B'),('B','C'),('C','D'),('D','A'),('A','E'),('E','B'),('B','D')])
    offsets={'A':(-12,8),'B':(3,-14),'C':(5,-8),'D':(5,5),'E':(-16,3),'M':(6,-9),'N':(-16,5),'H':(4,6),'K':(-14,-12)}
    if key=='q1sas':
        d.triangle(['A','B','M'],RED);d.triangle(['E','B','N'],BLUE)
        d.segment('E','N',BLUE,1.5,True)
        d.tick('B','A',1);d.tick('B','E',1);d.tick('B','M',2);d.tick('B','N',2)
        d.angle('B','A','M','45°',RED,14);d.angle('B','E','N','45°',BLUE,24)
    elif key=='q1equilateral':
        d.triangle(['B','M','N'],RED);d.segment('M','N',BLUE,1.8,True)
        for a,b in [('B','M'),('B','N'),('M','N')]:d.tick(a,b,1)
        d.angle('B','M','N','60°',RED,20)
    elif key in ['q1path','q1equal']:
        d.edges([('E','N'),('N','M'),('M','C')],color=RED,width=2)
        d.segment('E','C',BLUE,1.4,True)
        if key=='q1equal':
            d.segment('A','C',GRAY,.8,True)
            d.angle('C','E','B','15°',BLUE,13)
            d.angle('C','A','B','45°',GREEN,32)
            offsets['M']=(-2,8)
        else:d.segment('B','N',GRAY);d.segment('B','M',GRAY)
    else:
        d.edges([('E','H'),('E','K'),('K','B')],color=BLUE,width=1.4,dash=True)
        d.segment('E','C',RED,1.8);d.right('H','E','B');d.right('K','E','C')
        d.label_line('E','H','√3',(0,6));d.label_line('E','K','1',(-10,0));d.label_line('K','C','2+√3',(0,-20))
    d.labels(offsets)


def draw_q2(c,key,txt):
    s=math.sqrt(3)
    p={'A':(0.,0.),'B':(20.,0.),'C':(2.,0.),'D':(18.,0.),'P':(8.,0.),'Q':(10.,10*s),'E':(4.,4*s),'F':(14.,6*s),'M':(9.,5*s)}
    if key=='q2line':
        p={k:v for k,v in p.items() if k not in ['E','F']};p['H']=(5.,5*s)
    if key=='q2length':
        p={k:v for k,v in p.items() if k not in ['E','F','P','M']};p.update(K=(6.,5*s),L=(14.,5*s))
    d=Drawing(c,p,txt)
    d.edges([('A','B'),('A','Q'),('B','Q')])
    offsets={'A':(-12,-8),'B':(4,-8),'C':(-2,-17),'D':(-2,-17),'P':(-3,-14),'E':(-13,2),'F':(6,2),'M':(-14,6),'Q':(-4,9),'H':(-15,4),'K':(-4,-14),'L':(5,4)}
    if key in ['q2add','q2sides','q2sas']:
        d.edges([('A','E'),('E','P'),('P','A'),('B','F'),('F','P'),('P','B'),('E','F')],color=INK)
        d.edges([('E','Q'),('Q','F')],color=BLUE,width=1.4,dash=True)
        if key=='q2add':d.segment('Q','P',GREEN,1.4,True)
        if key=='q2sides':
            d.segment('E','P',RED,1.9);d.segment('F','Q',RED,1.9)
            d.tick('E','P',1);d.tick('F','Q',1)
        if key=='q2sas':
            d.triangle(['P','E','M'],RED);d.triangle(['Q','F','M'],BLUE)
            d.segment('P','Q',GREEN,1.3,True)
            d.tick('E','M',2);d.tick('F','M',2);d.tick('E','P',1);d.tick('F','Q',1)
    elif key=='q2line':
        d.triangle(['Q','A','P'],BLUE)
        d.segment('H','M',RED,2)
        d.tick('Q','H',1);d.tick('H','A',1);d.tick('Q','M',2);d.tick('M','P',2)
    else:
        d.triangle(['Q','C','D'],BLUE)
        d.segment('K','L',RED,2)
        d.tick('Q','K',1);d.tick('K','C',1);d.tick('Q','L',2);d.tick('L','D',2)
        d.label_line('C','D','16',(0,-33));d.label_line('K','L','8',(0,7))
    d.labels(offsets)


def draw_q3(c,key,txt):
    kind='square' if key.startswith('q3s') else ('general' if key.startswith('q3g') else 'outside')
    p=base_q3(kind)
    d=Drawing(c,p,txt)
    d.edges([('A','B'),('B','C'),('C','D'),('D','A'),('A','E'),('A','F'),('E','F')])
    d.segment('A','G',BLUE,1.4,True)
    d.segment('B','G',BLUE,1.4,True)
    if kind=='outside':d.edges([('C','E'),('D','F')],color=GRAY,dash=True)
    if key.endswith('1'):
        d.triangle(['A','B','G'],RED);d.triangle(['A','D','F'],BLUE)
        d.tick('A','B',1);d.tick('A','D',1);d.tick('B','G',2);d.tick('D','F',2)
    else:
        d.triangle(['G','A','E'],RED);d.triangle(['F','A','E'],BLUE)
        d.tick('A','G',1);d.tick('A','F',1);d.tick('A','E',2)
        d.segment('G','E',GREEN,2);d.segment('E','F',GREEN,2)
    offsets={'A':(-10,9),'B':(-14,-14),'C':(5,-10),'D':(5,7),'E':(-4,-16),'F':(5,4),'G':(-11,-10)}
    if kind=='outside':offsets.update(G=(-2,6),E=(5,-8),F=(-2,9),D=(-15,3))
    d.labels(offsets)


def draw_lemma(c,key,txt):
    p={'T':(0.,4.),'U':(-3.,0.),'V':(3.,0.),'X':(-1.5,2.),'Y':(1.5,2.),'Z':(4.5,2.)}
    d=Drawing(c,p,txt)
    d.edges([('T','U'),('U','V'),('V','T'),('X','Y')])
    d.segment('Y','Z',BLUE,1.4,True)
    if key=='lemma1':
        d.triangle(['T','Y','X'],RED);d.triangle(['V','Y','Z'],BLUE)
        d.tick('T','Y',1);d.tick('Y','V',1);d.tick('X','Y',2);d.tick('Y','Z',2)
    elif key=='lemma2':
        d.triangle(['X','U','Z'],RED);d.triangle(['V','Z','U'],BLUE)
        d.segment('U','Z',GREEN,1.6,True);d.tick('U','X',1);d.tick('V','Z',1)
    else:
        d.segment('X','Y',RED,2);d.segment('U','V',RED,2)
        d.segment('V','Z',GRAY,.8,True)
        d.tick('X','Y',1);d.tick('Y','Z',1)
    d.labels({'T':(-3,9),'U':(-13,-10),'V':(4,-13),'X':(-14,2),'Y':(-2,7),'Z':(5,4)})


def step(title, figure, action, proof, result):
    return {'title':title,'figure':figure,'action':action,'proof':proof,'result':result}


Q1=[
 step('第1步：把AM换成EN','q1sas','连EN，突出△ABM和△EBN。',
 '①BA=BE（等边△ABE）。\n②BM=BN（旋转得到）。\n③∠ABM=∠EBN=45°。\n所以△ABM≌△EBN（SAS）。对应顶点是A与E、B与B、M与N。', '对应边：AM=EN。'),
 step('第2步：把BM换成MN','q1equilateral','连MN，单独看△BMN。',
 '①BM=BN。\n②∠MBN=60°。\n等腰三角形的两个底角均为(180°-60°)/2=60°。\n因此△BMN是等边三角形。', 'BM=MN；此时三段距离成为EN、NM、MC。'),
 step('第3步：把长度和写成折线路径','q1path','蓝色虚线EC是固定线段；红色EN、NM、MC是折线。',
 'AM+BM+CM\n=EN+NM+MC\n≥EC。\n依据：两点之间线段最短。这里只找到了下界，还必须证明等号能取到。若M=B，三段和为4，且EC<BE+BC=4。', '任何允许的M，都有AM+BM+CM≥EC。'),
 step('第4步：画出取得最小值的位置','q1equal','取M为EC与BD的交点，再得到N；补画AC确认取点范围。',
 'BE=BC，∠EBC=150°，所以∠ECB=15°。\n∠BCA=45°，CE进入△ABC内部，故与BD相交于BD内部。\nBD与BC夹45°；由等边△BMN，MN与BC夹15°，所以MN∥EC。\n当M在EC上时，E、N、M、C依次共线，等号成立。', '等号可取得；最小值为CE，M=EC与BD的交点。'),
 step('第5步：数值计算（未学勾股可跳过）','q1number','作EH⊥AB于H；作EK⊥BC所在直线于K。',
 '等边△ABE边长2，H为AB中点，所以BH=1。\nEH²=2²-1²=3，EH=√3。\nEK=BH=1，KC=KB+BC=√3+2。\nCE²=EK²+KC²=8+4√3。', 'CE=√6+√2。主要诊断评分可只要求“最小值为CE”。'),
]
Q2=[
 step('第1步：补一个不会动的等边三角形','q2add','在AB同侧作等边△ABQ，Q是固定点；连QP。',
 '∠EAP=∠QAB=60°，所以A、E、Q共线。\n同理B、F、Q共线。\nE、F随P运动，但Q不动。\n后面要把M的运动转化成“固定点Q与动点P之间的中点运动”。', '固定辅助点：Q；E在AQ上，F在BQ上。'),
 step('第2步：找到一组等长且平行的边','q2sides','突出EP和FQ；二者都沿等边三角形的同一边方向。',
 '由两个等边三角形的60°角，EP∥BQ，因此EP∥FQ。\nEP=AP，BF=BP，BQ=AB。\nFQ=BQ-BF=AB-BP=AP。', 'EP=FQ，EP∥FQ。'),
 step('第3步：证明M也恰好是QP中点','q2sas','连PM、QM，突出△PEM和△QFM。',
 '①EM=FM（M为EF中点）。\n②EP=FQ（第2步）。\n③∠PEM=∠QFM（平行线）。\n所以△PEM≌△QFM（SAS），PM=QM。P、Q在EF两侧，且∠EMP=∠FMQ、EM与FM反向共线，故MP、MQ也反向共线。', 'Q、M、P共线且QM=MP，M是QP中点。'),
 step('第4步：证明M走的是固定直线','q2line','取QA中点H，连HM；单独看△QAP。',
 'H为QA中点，M为QP中点。\n由中点连线性质，HM∥AP。\nA、P都在固定直线AB上，所以HM始终平行AB。H也固定，因此M始终在过H且平行AB的同一条直线上。', 'M沿固定直线运动，不是曲线；未学中点性质先看附录。'),
 step('第5步：定位起点终点，再算路径长','q2length','取QC、QD的中点K、L，连KL。',
 'P=C时，M=K；P=D时，M=L。\n在△QCD中，KL是中点连线。\nCD=AB-AC-BD=20-2-2=16。\nKL=CD/2=8。P单向从C到D，M也单向从K到L。', 'M的运动路径是线段KL，路径长8。'),
]
Q3=[
 step('（1）第1步：在B外侧复制DF','q3s1','延长CB到G，使BG=DF，连AG。',
 '正方形给AB=AD，∠ABG=∠ADF=90°。\n再加BG=DF，得到\n△ABG≌△ADF（SAS）。\n所以AG=AF，∠BAG=∠DAF。', '第一次全等得到：AG=AF，以及一组对应角相等。'),
 step('（1）第2步：第二次全等，得到加法关系','q3s2','突出△GAE和△FAE；红绿线段GE由GB与BE组成。',
 '∠GAE=∠BAG+∠BAE\n=∠DAF+∠BAE\n=90°-45°=45°=∠EAF。\nAG=AF，AE公共，所以△GAE≌△FAE（SAS），GE=FE。\nG、B、E依次共线，GE=GB+BE。', 'EF=GE=GB+BE=DF+BE。'),
 step('（2）第1步：一般四边形中仍复制DF','q3g1','仍延长CB到G，使BG=DF，连AG。',
 'AB=AD。\n∠ABG=180°-∠ABC。\n已知∠ABC+∠ADC=180°，所以∠ABG=∠ADC=∠ADF。\n结合BG=DF，得到△ABG≌△ADF（SAS）。', '仍有AG=AF，∠BAG=∠DAF；不需要四个角都为90°。'),
 step('（2）第2步：用“半角”替代45°','q3g2','突出第二组全等三角形；半角条件用在A处。',
 '∠GAE=∠BAG+∠BAE\n=∠DAF+∠BAE\n=∠BAD-∠EAF\n=∠EAF（半角条件）。\n由AG=AF、AE公共，△GAE≌△FAE（SAS），GE=EF。\n依然G、B、E依次共线。', '结论仍成立：EF=BE+DF。'),
 step('（3）第1步：G这次放在B、E之间','q3o1','在射线BC上取G，使BG=DF；连AG。不是延长CB。',
 'F在CD越过D的延长线上。\n∠ADF=180°-∠ADC=∠ABC=∠ABG。\nAB=AD，BG=DF，所以△ABG≌△ADF（SAS）。\n得到AG=AF，∠BAG=∠DAF。', 'G改放在B朝C的方向；第一组全等仍成立。'),
 step('（3）第2步：相等的是GE，最终要做减法','q3o2','突出△GAE和△FAE，核对B、G、E的顺序。',
 '∠BAE=∠EAF+∠DAF，所以∠GAE=∠BAE-∠BAG=∠EAF。\nAG=AF、AE公共，得到△GAE≌△FAE（SAS），GE=EF。\n0<∠BAG<∠BAE，所以AG在∠BAE内部，G在BE内部。\nGE=BE-BG=BE-DF。', 'BE=DF+EF（与前两问的加法关系不同）。'),
]
Q4=[
 step('（1）第1步：PM是哪一个三角形的中点连线？','q4a1','先只突出△DCE，PM用红色画出。',
 '题目说P是DC中点，M是DE中点。\n所以DP=PC，DM=ME。\nPM连接△DCE的两边中点。\n由中点连线性质：\nPM∥CE，PM=CE/2。\n若学生未学此性质，先讲附录配图证明。', '第一条关系：PM∥CE，PM=CE/2。'),
 step('（1）第2步：PN属于另一个三角形','q4a2','现在改为突出△DCB，PN用红色画出。',
 'P仍是DC中点，但N是BC中点。\n所以DP=PC，BN=NC。\nPN连接△DCB的两边中点。\n因此PN∥DB，PN=DB/2。\n注意：PM与PN分别属于两个不同的大三角形。', '第二条关系：PN∥DB，PN=DB/2。'),
 step('（1）第3步：先用初始位置得出答案','q4a3','观察未旋转的图：D在AB上，E在AC上。',
 'BD=AB-AD，CE=AC-AE。\nAB=AC，AD=AE，所以BD=CE。\nAB⊥AC，因此BD⊥CE。\n结合前两步，PM=CE/2=BD/2=PN；PM∥CE、PN∥BD，所以PM⊥PN。', '第（1）问：PM=PN，PM⊥PN。'),
 step('（2）第1步：连接BD、CE，证明等长','q4sas','连BD、CE；红色△BAD与蓝色△CAE作对照。',
 '图2中∠BAC=∠DAE=90°，且D在∠BAC内部。\n∠BAD=90°-∠DAC。\n∠CAE=90°-∠DAC。\n所以∠BAD=∠CAE。\nAB=AC，AD=AE，故△BAD≌△CAE（SAS）。', 'BD=CE；对应顶点是B与C、A与A、D与E。'),
 step('（2）第2步：用角度一步步证明垂直','q4angles','延长BD与CE交于T；突出△BCT，令x=∠ABD。',
 '由刚才全等，∠ACE=∠ABD=x。\n等腰直角△ABC中，∠ABC=∠BCA=45°。\n∠CBT=45°-x；\n∠BCT=45°+x。\n所以∠BTC=180°-(45°-x)-(45°+x)=90°。\n这里使用图2的点与角位置。', 'BD⊥CE。“两边相等”本身不能直接推出垂直。'),
 step('（2）第3步：旋转后PM仍是中点连线','q4pm','旋转图中单独突出△DCE；突出DP、PC与DM、ME。',
 'D、E的位置改变了，但P始终是DC中点，M始终是DE中点。\n因此在新位置的△DCE中，仍然有\nPM∥CE，PM=CE/2。\n不需要判断PM在图上看起来是水平还是斜着。', '旋转后依然：PM∥CE，PM=CE/2。'),
 step('（2）第4步：旋转后PN仍属于△DCB','q4pn','只突出△DCB及红色PN，确认P、N分别在哪两边。',
 'P是DC中点，N是CB中点。\n所以在新位置的△DCB中，\nPN∥DB，PN=DB/2。\n两条小线段的长度与方向，分别由CE和DB决定。', '旋转后依然：PN∥DB，PN=DB/2。'),
 step('（2）第5步：合并等长、垂直两个结论','q4result','突出小三角形PMN，在P处标直角。',
 '因为BD=CE，\nPM=CE/2=BD/2=PN。\n因为BD⊥CE，且PN∥BD、PM∥CE，\n所以PM⊥PN。\nPM与PN相等并互相垂直，故△PMN是等腰直角三角形。', '第（2）问：△PMN为等腰直角三角形，直角顶点为P。'),
 step('（3）第1步：先把面积写成BD的式子','q4area','突出两条直角边PM、PN，各标为BD/2。',
 '自由旋转时，绕A逆时针转90°始终把B对应到C、D对应到E，所以BD与CE始终等长且垂直。\n于是PM=PN=BD/2。\nS△PMN=(1/2)×PM×PN\n=(1/2)×(BD/2)×(BD/2)\n=BD²/8。', '面积只随BD的长度改变；BD越长，面积越大。'),
 step('（3）第2步：BD最多多长？','q4bound','单独画B、A、D及AB=10、AD=4。',
 '从B经过A到D是折线，BD是直接连接B、D的线段。\n因此BD≤BA+AD=10+4=14。\n一般不共线的位置严格小于14。\n要等于14，必须让B、A、D依次共线，也就是A在B与D之间。', 'BD的上界为14；必须检查这个位置能否通过旋转得到。'),
 step('（3）第3步：画出达到最大值的完整位置','q4max','把AD旋转到AB的反向延长线上；重画E、M、P、N。',
 '自由旋转允许D、A、B依次共线。\n此时BD=DA+AB=4+10=14。\nCE=14，PM=PN=14/2=7。\nS△PMN=(1/2)×7×7=49/2。\n图中小三角形仍然是等腰直角三角形，没有退化。', '第（3）问：最大面积49/2；D在AB的反向延长线上时取得。'),
]
LEMMA=[
 step('引理第1步：延长中点连线，制造第一组全等','lemma1','△TUV中X、Y为TU、TV中点。延长XY至Z，使YZ=YX。',
 '①TY=YV（Y是中点）。\n②YX=YZ（复制长度）。\n③∠TYX=∠VYZ（对顶角）。\n所以△TYX≌△VYZ（SAS）。\n对应边VZ=TX=UX；对应角∠YTX=∠YVZ，故VZ∥UX。', 'VZ=UX，VZ∥UX。'),
 step('引理第2步：连UZ，制造第二组全等','lemma2','连UZ，突出△XUZ与△VZU。',
 '①UX=VZ（第1步）。\n②UZ是公共边。\n③∠XUZ=∠VZU（UX∥VZ，内错角相等）。\n所以△XUZ≌△VZU（SAS）。\n对应边XZ=VU；对应角∠XZU=∠VUZ，故XZ∥VU。', 'XZ=UV，XZ∥UV。'),
 step('引理第3步：从整段XZ回到半段XY','lemma3','突出原来的XY、底边UV，确认X、Y、Z的顺序。',
 'X、Y、Z依次共线，XY=YZ。\n所以XZ=XY+YZ=2XY。\n第2步已经证明XZ=UV，因此XY=UV/2。\nXY与XZ在同一直线上，故XY∥UV。', '中点连线性质：XY∥UV，XY=UV/2。两次SAS即可证明。'),
]
PARTS=[('第1题：旋转与三段距离最小值',Q1),('第2题：等边动点与中点轨迹',Q2),('第3题：基本图、一般图和延长线变式',Q3),('第4题：逐步图解',Q4),('附录：中点连线性质的全等证明',LEMMA)]


def figures(root, txt):
    out=root/'teaching_figures';out.mkdir(exist_ok=True)
    for _,steps in PARTS:
        for s in steps:
            key=s['figure'];buf=io.BytesIO()
            c=canvas.Canvas(buf,pagesize=(260,220))
            c.setTitle(s['title'])
            if key.startswith('q1'):draw_q1(c,key,txt)
            elif key.startswith('q2'):draw_q2(c,key,txt)
            elif key.startswith('q3'):draw_q3(c,key,txt)
            elif key.startswith('q4'):draw_q4(c,key,txt)
            else:draw_lemma(c,key,txt)
            c.save()
            doc=fitz.open(stream=buf.getvalue(),filetype='pdf')
            (out/(key+'.pdf')).write_bytes(buf.getvalue())
            doc[0].get_pixmap(matrix=fitz.Matrix(3,3)).save(out/(key+'.png'))


def diagram_at(c,key,x,y,w,h,txt):
    c.saveState();c.translate(x,y);c.scale(w/260,h/220)
    if key.startswith('q1'):draw_q1(c,key,txt)
    elif key.startswith('q2'):draw_q2(c,key,txt)
    elif key.startswith('q3'):draw_q3(c,key,txt)
    elif key.startswith('q4'):draw_q4(c,key,txt)
    else:draw_lemma(c,key,txt)
    c.restoreState()


def block(c,s,top,txt):
    txt(c,s['title'],18*MM,top*MM,size=11.5,lead=16)
    diagram_at(c,s['figure'],18*MM,(top-81)*MM,82*MM,69*MM,txt)
    body='作图：'+s['action']+'\n\n'+s['proof']
    end=txt(c,body,107*MM,(top-11)*MM,width=85*MM,size=10,lead=14)
    if end<(top-85)*MM:
        raise ValueError('Step text exceeds allotted area: '+s['figure'])
    c.setStrokeColorRGB(*BLUE);c.setLineWidth(.6)
    c.line(18*MM,(top-87)*MM,192*MM,(top-87)*MM)
    txt(c,'本步结论：'+s['result'],18*MM,(top-94)*MM,size=10.5,lead=15)


def pdf(root,guide,txt,header,only_q4=False):
    name='question4_step_by_step.pdf' if only_q4 else 'selected_teacher.pdf'
    c=canvas.Canvas(str(root/name),pagesize=PAGE)
    c.setTitle('第4题逐步图解' if only_q4 else '现有题选编：逐步配图教师答案与两小时教案')
    page=1
    def start(title):
        header(c,'第4题：逐步图解' if only_q4 else '教师答案：逐步配图与辅助线','图左看关系，图右写证明｜彩色突出本步对象｜虚线为辅助或延长线',page,revision='2026-10-13')
        # Cover and teaching pages identify the revised date independently of source date.
        txt(c,title+'｜2026-10-13更新',18*MM,255*MM,size=10,lead=15)
    if not only_q4:
        start('使用说明与两小时安排')
        y=241*MM
        sections=[('怎么读这版','每个证明步骤都配一张独立重绘的数学图。先按“作图”添加辅助线，再对照右侧逐条证明，最后看“本步结论”。同一种刻痕表示等长线段，同样的弧线标记表示相等的角；彩色三角形表示当前比较的对象。原题条件与学生卷不变，教学图只表示一个合法位置，不按给定数值比例绘制。'),*guide[0][1][2:5],('页码导航','第1题第2—4页；第2题第5—7页；第3题第8—10页；第4题第11—16页；中点引理第17—18页；评分及回练第19页。第4题另有6页独立文件question4_step_by_step.pdf。')]
        for title,body in sections:
            # Replace old fixed-page references after expanding the teaching guide.
            body=body.replace('第6页','附录')
            y=txt(c,title,18*MM,y,size=11,lead=17)
            y=txt(c,body,18*MM,y-2,size=10.5,lead=16)-11
        if y<24*MM:raise ValueError('Cover overflow')
        c.showPage();page+=1
    parts=[PARTS[3]] if only_q4 else PARTS
    for title,steps in parts:
        for i in range(0,len(steps),2):
            start(title)
            block(c,steps[i],242,txt)
            if i+1<len(steps):block(c,steps[i+1],132,txt)
            c.showPage();page+=1
    if not only_q4:
        start('评分、讲解选择与即时回练')
        y=241*MM
        sections=[]
        for idx in range(1,5):
            for title,body in guide[idx][1]:
                if '评分' in title:sections.append((f'第{idx}题：'+title,body))
        sections.extend(guide[-1][1][1:3])
        sections.append(guide[-1][1][-1])
        for title,body in sections:
            y=txt(c,title,18*MM,y,size=10.5,lead=16)
            y=txt(c,body,18*MM,y-2,size=9.7,lead=14)-8
        if y<23*MM:raise ValueError('Scoring page overflow')
        c.showPage()
    c.save()


def word(root,guide,word_base,only_q4=False):
    name='question4_step_by_step.docx' if only_q4 else 'selected_teacher.docx'
    d=word_base('第4题逐步图解' if only_q4 else '现有题选编：逐步配图教师答案与两小时教案')
    d.sections[0].footer.paragraphs[0].text='逐步教学图与辅助线｜2026-10-13更新｜原题出处见README'
    d.styles['Normal'].font.size=Pt(10)
    if not only_q4:
        d.add_heading('逐步配图教师答案与两小时教案',0)
        d.add_paragraph('每一步均配数学图、作图指令、分条证明与本步结论。原题条件不变；教学图不按数值比例绘制。')
        for title,body in guide[0][1][2:5]:
            d.add_heading(title,2);d.add_paragraph(body.replace('第6页','附录'))
    parts=[PARTS[3]] if only_q4 else PARTS
    first=True
    for title,steps in parts:
        for s in steps:
            if not first or not only_q4:d.add_page_break()
            first=False
            d.add_heading(title,0)
            d.add_heading(s['title'],1)
            d.add_paragraph('作图：'+s['action'])
            d.add_picture(str(root/'teaching_figures'/(s['figure']+'.png')),width=Cm(10.5))
            d.add_paragraph(s['proof'])
            d.add_paragraph('本步结论：'+s['result'])
    if not only_q4:
        d.add_page_break();d.add_heading('评分与课堂回练',0)
        for idx in range(1,5):
            for title,body in guide[idx][1]:
                if '评分' in title:d.add_heading(f'第{idx}题：'+title,2);d.add_paragraph(body)
        for title,body in guide[-1][1][1:]:
            d.add_heading(title,2);d.add_paragraph(body)
    d.save(root/name)


def markdown(root,guide,only_q4=False):
    name='question4_step_by_step.md' if only_q4 else 'selected_teacher.md'
    text='# '+('第4题：逐步配图与步骤答案' if only_q4 else '教师答案：逐步配图与辅助线')+'\n\n2026-10-13更新。原题出处与原图见[README](README.md)。教学图为独立重绘，辅助线与结论按步骤呈现。\n\n'
    if not only_q4:
        for title,body in guide[0][1]:
            text+='## '+title+'\n\n'+body.replace('第6页','附录')+'\n\n'
    parts=[PARTS[3]] if only_q4 else PARTS
    for title,steps in parts:
        text+='## '+title+'\n\n'
        for s in steps:
            text+='### '+s['title']+'\n\n'
            text+='![该步的图形与辅助线](teaching_figures/'+s['figure']+'.png)\n\n'
            text+='**作图：** '+s['action']+'\n\n'+s['proof'].replace('\n','\n\n')+'\n\n'
            text+='**本步结论：** '+s['result']+'\n\n'
    if not only_q4:
        text+='## 评分与课堂回练\n\n'
        for idx in range(1,5):
            for title,body in guide[idx][1]:
                if '评分' in title:text+=f'### 第{idx}题：{title}\n\n{body}\n\n'
        for title,body in guide[-1][1][1:]:text+=f'### {title}\n\n{body}\n\n'
    (root/name).write_text(text.rstrip()+'\n',encoding='utf-8')
